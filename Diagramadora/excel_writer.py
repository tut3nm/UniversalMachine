"""
Genera el archivo Excel (.xlsx) a partir del resultado del parseo (ver parser.py).

Estructura del Excel:
  - Hoja "Resumen": datos generales del archivo + una tabla con cada configuracion
    de tolerancias encontrada (puede haber mas de una si el equipo fue recalibrado
    con el tiempo) + totales OK/NOK.
  - Hoja "Mediciones": una fila por cada ensayo individual, con los valores medidos,
    el resultado de cada control y un resultado general, con formato condicional
    (verde=OK, rojo=NOK) y autofiltro para que el ingeniero pueda filtrar/ordenar.
"""
from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from parser import ParseResult, ConfigBlock, Measurement

# ---------------------------------------------------------------------------
# Estilos (se crean una sola vez y se reutilizan en todas las celdas)
# ---------------------------------------------------------------------------

TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E78")
VALUE_FONT = Font(name="Calibri", size=10)
NOTE_FONT = Font(name="Calibri", size=9, italic=True, color="808080")

HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

DATA_FONT = Font(name="Calibri", size=10)
CENTER = Alignment(horizontal="center", vertical="center")

THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CFG_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTALS_FILL = PatternFill("solid", fgColor="FFF2CC")

OK_FILL = PatternFill("solid", fgColor="C6EFCE")
OK_FONT = Font(color="006100")
NOK_FILL = PatternFill("solid", fgColor="FFC7CE")
NOK_FONT = Font(color="9C0006")


def _num(v):
    if v is None:
        return None
    return int(v) if float(v).is_integer() else v


def _autosize(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------------------
# Hoja Resumen
# ---------------------------------------------------------------------------


def _write_resumen(wb: Workbook, result: ParseResult, source_path: str) -> None:
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False

    row = 1
    ws.cell(row=row, column=1, value="Ensayos de Fuerza-Velocidad de Amortiguadores").font = TITLE_FONT
    row += 1
    ws.cell(row=row, column=1, value=f"Archivo origen: {os.path.basename(source_path)}").font = SUBTITLE_FONT
    row += 1
    ws.cell(
        row=row, column=1,
        value=f"Excel generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
    ).font = SUBTITLE_FONT
    row += 2

    total = len(result.measurements)
    total_ok = sum(1 for m in result.measurements if m.resultado_general)
    total_nok = total - total_ok
    con_fecha = [m for m in result.measurements if m.dt is not None]
    primera = con_fecha[0].dt if con_fecha else None
    ultima = con_fecha[-1].dt if con_fecha else None

    ws.cell(row=row, column=1, value="TOTALES GENERALES").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    for c in range(2, 6):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1

    pares = [
        ("Cantidad de mediciones", total),
        ("Mediciones OK", total_ok),
        ("Mediciones NOK", total_nok),
        ("% OK", f"{(100 * total_ok / total):.2f}%" if total else "-"),
        ("Primera medición (según orden del archivo)", primera.strftime("%d/%m/%Y %H:%M:%S") if primera else "-"),
        ("Última medición (según orden del archivo)", ultima.strftime("%d/%m/%Y %H:%M:%S") if ultima else "-"),
        ("Configuraciones de tolerancia impresas por el equipo", len(result.configs)),
    ]
    for label, value in pares:
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        cell = ws.cell(row=row, column=2, value=value)
        cell.font = VALUE_FONT
        cell.fill = TOTALS_FILL
        row += 1

    row += 1
    ws.cell(
        row=row, column=1,
        value=(
            "Nota: el criterio de aceptacion de Extendido/Comprimido/Fuerza es "
            "limite_inferior <= valor < limite_superior (tocar el limite superior "
            "exacto ya se considera fuera de tolerancia). Este criterio fue "
            "verificado contra las banderas que imprime el propio equipo (0 "
            "diferencias sobre la totalidad de las mediciones del archivo)."
        ),
    ).font = NOTE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    ws.row_dimensions[row].height = 28
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    row += 1
    ws.cell(
        row=row, column=1,
        value=(
            "Nota: cada medicion trae ademas banderas 'Y' (extendido/comprimido "
            "dentro de tolerancia, coincide con las columnas Resultado de esta "
            "planilla), 'V' y 'F' (fuerza a velocidad fija) generadas por el propio "
            "equipo. La bandera 'V' no esta documentada en el archivo de "
            "anotaciones entregado: se muestra tal cual la imprime el equipo, sin "
            "reinterpretarla, y SI se tiene en cuenta para el Resultado General."
        ),
    ).font = NOTE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
    ws.row_dimensions[row].height = 40
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    row += 2

    if result.date_anomalies:
        ws.cell(row=row, column=1, value="ANOMALÍAS DE FECHA DETECTADAS EN EL EQUIPO").font = SECTION_FONT
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor="C00000")
        for c in range(2, 8):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="C00000")
        row += 1
        ws.cell(
            row=row, column=1,
            value=(
                "El reloj del equipo retrocedió en los siguientes tramos (posible "
                "batería de reloj agotada o fecha mal configurada). Las mediciones "
                "en sí siguen siendo válidas: solo la fecha/hora registrada es "
                "incorrecta en esas filas. No se modificó ningún dato: se informa "
                "tal cual viene del archivo original para que se pueda corregir "
                "manualmente si corresponde."
            ),
        ).font = NOTE_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        ws.row_dimensions[row].height = 40
        ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

        anom_headers = ["Fila desde", "Fila hasta", "Cant. filas", "Fecha mostrada (desde)", "Fecha mostrada (hasta)", "Última fecha válida previa"]
        for c, h in enumerate(anom_headers, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER
        row += 1
        for a in result.date_anomalies:
            values = [
                a.row_ini, a.row_fin, a.row_fin - a.row_ini + 1,
                a.dt_ini.strftime("%d/%m/%Y %H:%M:%S"),
                a.dt_fin.strftime("%d/%m/%Y %H:%M:%S"),
                a.dt_referencia.strftime("%d/%m/%Y %H:%M:%S"),
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.font = DATA_FONT
                cell.border = BORDER
                cell.alignment = CENTER
            row += 1
        row += 1

    ws.cell(row=row, column=1, value="CONFIGURACIONES DE ENSAYO / TOLERANCIAS").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    for c in range(2, 12):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1

    headers = [
        "Config.", "Vigente desde", "Archivo / Programa (ARCHIVO)", "Carrera (mm)",
        "Velocidades ensayo (mm/s)",
        "Tolerancia Extendido (kgf) [inf, sup)",
        "Tolerancia Comprimido (kgf) [inf, sup)",
        "CPCAVIT (crudo, sin documentar)",
        "Fuerza a vel. fija: veloc. / rango (kgf) [inf, sup)",
        "Cant. mediciones", "OK", "NOK", "% OK",
    ]
    header_row = row
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    row += 1

    for cfg in result.configs:
        if cfg.n_measurements == 0:
            # bloque de cabecera duplicado que nunca llego a tener mediciones propias
            # (ej. el par inicial "Prod"+"COP" con la misma marca de tiempo)
            continue
        ext_tol = "; ".join(
            f"V{_num(v)}: [{_num(lo)}, {_num(hi)})"
            for v, lo, hi in zip(cfg.velocities, cfg.ext_inf, cfg.ext_sup)
        )
        comp_tol = "; ".join(
            f"V{_num(v)}: [{_num(lo)}, {_num(hi)})"
            for v, lo, hi in zip(cfg.velocities, cfg.comp_inf, cfg.comp_sup)
        )
        fza = (
            f"V{_num(cfg.fv_speed)}: [{_num(cfg.fv_min)}, {_num(cfg.fv_max)})"
            if cfg.fv_speed is not None else "-"
        )
        values = [
            cfg.index,
            cfg.effective_from.strftime("%d/%m/%Y %H:%M:%S") if cfg.effective_from else "-",
            cfg.archivo,
            _num(cfg.carrera),
            ", ".join(str(_num(v)) for v in cfg.velocities),
            ext_tol,
            comp_tol,
            ", ".join(str(_num(v)) for v in cfg.cpcavit),
            fza,
            cfg.n_measurements,
            cfg.n_ok,
            cfg.n_nok,
            f"{(100 * cfg.n_ok / cfg.n_measurements):.1f}%" if cfg.n_measurements else "-",
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = DATA_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.column_letter == "A":
                cell.fill = CFG_HEADER_FILL
                cell.alignment = CENTER
        row += 1

    _autosize(ws, {
        "A": 9, "B": 19, "C": 26, "D": 12, "E": 18, "F": 26, "G": 26,
        "H": 22, "I": 26, "J": 14, "K": 9, "L": 9, "M": 9,
    })


# ---------------------------------------------------------------------------
# Hoja Mediciones
# ---------------------------------------------------------------------------


def _write_mediciones(
    wb: Workbook,
    result: ParseResult,
    progress_callback=None,
) -> None:
    ws = wb.create_sheet("Mediciones")

    max_slots = max((len(c.velocities) for c in result.configs), default=0)

    headers = ["N°", "Fecha", "Hora", "Configuración"]
    for i in range(max_slots):
        n = i + 1
        headers += [
            f"V{n} Velocidad (mm/s)",
            f"V{n} Extendido (kgf)",
            f"V{n} Result. Extendido",
            f"V{n} Comprimido (kgf)",
            f"V{n} Result. Comprimido",
            f"V{n} Flags equipo (Y/V)",
        ]
    headers += [
        "Fza.Fija Velocidad (mm/s)",
        "Fza.Fija Medida (kgf)",
        "Fza.Fija Resultado",
        "Fza.Fija Flag equipo (F)",
        "Resultado General",
    ]

    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    n = len(result.measurements)
    report_every = max(1, n // 100)

    for r, m in enumerate(result.measurements, start=2):
        if progress_callback and r % report_every == 0:
            progress_callback(r - 2, n)

        col = 1
        ws.cell(row=r, column=col, value=m.row_num); col += 1

        fecha_cell = ws.cell(row=r, column=col, value=m.dt.date() if m.dt else None)
        fecha_cell.number_format = "dd/mm/yyyy"
        col += 1
        hora_cell = ws.cell(row=r, column=col, value=m.dt.time() if m.dt else None)
        hora_cell.number_format = "hh:mm:ss"
        col += 1

        cfg = m.config
        cfg_txt = f"{cfg.index} ({cfg.effective_from.strftime('%d/%m/%Y')})" if cfg and cfg.effective_from else (str(cfg.index) if cfg else "-")
        ws.cell(row=r, column=col, value=cfg_txt); col += 1

        ec_by_slot = list(m.ec) + [None] * max(0, max_slots - len(m.ec))
        for ec in ec_by_slot[:max_slots]:
            if ec is None:
                col += 6
                continue
            ws.cell(row=r, column=col, value=_num(ec.velocidad)); col += 1
            ws.cell(row=r, column=col, value=_num(ec.extendido)); col += 1
            ws.cell(row=r, column=col, value=("OK" if ec.ok_ext else ("NOK" if ec.ok_ext is False else "-"))); col += 1
            ws.cell(row=r, column=col, value=_num(ec.comprimido)); col += 1
            ws.cell(row=r, column=col, value=("OK" if ec.ok_comp else ("NOK" if ec.ok_comp is False else "-"))); col += 1
            ws.cell(row=r, column=col, value=f"Y{ec.flag_y} V{ec.flag_v}"); col += 1

        ws.cell(row=r, column=col, value=_num(m.v70_speed)); col += 1
        ws.cell(row=r, column=col, value=_num(m.v70_force)); col += 1
        ws.cell(row=r, column=col, value=("OK" if m.v70_ok else ("NOK" if m.v70_ok is False else "-"))); col += 1
        ws.cell(row=r, column=col, value=(f"F{m.v70_flag}" if m.v70_flag else "-")); col += 1
        ws.cell(row=r, column=col, value=("OK" if m.resultado_general else "NOK")); col += 1

    last_row = n + 1

    # Formato condicional (reglas, no celda por celda -> rapido de generar y de abrir)
    ok_cols = []
    for i in range(max_slots):
        base = 5 + i * 6
        ok_cols.append(base + 2)  # columna "V{n} Result. Extendido"
        ok_cols.append(base + 4)  # columna "V{n} Result. Comprimido"
    ok_cols.append(4 + max_slots * 6 + 3)  # Fza.Fija Resultado
    general_col = 4 + max_slots * 6 + 5     # Resultado General
    ok_cols.append(general_col)

    for col_idx in ok_cols:
        letter = get_column_letter(col_idx)
        rng = f"{letter}2:{letter}{last_row}"
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"OK"'], fill=OK_FILL, font=OK_FONT)
        )
        ws.conditional_formatting.add(
            rng, CellIsRule(operator="equal", formula=['"NOK"'], fill=NOK_FILL, font=NOK_FONT)
        )

    widths = {"A": 8, "B": 12, "C": 10, "D": 16}
    for i in range(max_slots):
        base_letter_idx = 5 + i * 6
        for offset, w in zip(range(6), [10, 12, 13, 12, 13, 13]):
            widths[get_column_letter(base_letter_idx + offset)] = w
    fza_base = 4 + max_slots * 6 + 1
    for offset, w in zip(range(5), [10, 12, 12, 12, 13]):
        widths[get_column_letter(fza_base + offset)] = w
    _autosize(ws, widths)

    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).border = BORDER

    if progress_callback:
        progress_callback(n, n)


# ---------------------------------------------------------------------------
# API publica
# ---------------------------------------------------------------------------


def write_excel(
    result: ParseResult,
    output_path: str,
    source_path: str,
    progress_callback=None,
) -> None:
    wb = Workbook()
    _write_resumen(wb, result, source_path)
    _write_mediciones(wb, result, progress_callback=progress_callback)
    wb.active = 0
    wb.save(output_path)
