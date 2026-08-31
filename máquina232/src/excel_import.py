"""
excel_import.py
================
Lectura de archivos Excel (.xlsx/.xlsm) y CSV de formato variable para la
función "Importar cambios" del Configurador de Parámetros de Planta.

Como cada archivo puede traer las columnas en distinto orden o con
distintos encabezados, este módulo separa la lectura en dos pasos:

  1. `abrir_archivo_importacion` / `sheet_names` / `read_headers` /
     `guess_mapping_generic`: exploran el archivo para que la interfaz le
     pida al usuario qué columna corresponde a cada campo del perfil.
  2. `read_rows_generic`: una vez que el usuario confirmó el mapeo,
     recorre los datos usando esas columnas.
"""

from __future__ import annotations

import csv
import difflib
import unicodedata

# openpyxl es opcional a nivel de módulo (Nivel 4.6: importar desde CSV no
# debería requerirlo). Solo hace falta para abrir un .xlsx/.xlsm de
# verdad — open_workbook() lo importa recién ahí y da un error claro si
# falta, en vez de tumbar todo este módulo (y con él, la importación
# desde CSV) al arrancar la app.
try:
    import openpyxl
except ImportError:
    openpyxl = None


def _strip_accents(text: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", text)
                   if not unicodedata.combining(c))


def open_workbook(path: str):
    """Abre el libro en modo solo-lectura, con fórmulas resueltas a valor."""
    if openpyxl is None:
        raise RuntimeError(
            "Para importar desde Excel hace falta el paquete 'openpyxl', que "
            "no está instalado. Instalalo con: py -m pip install openpyxl "
            "(o importá desde un archivo .csv, que no lo necesita).")
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


# --- Importación desde CSV (Nivel 4.6) ---------------------------------------
class _HojaCSV:
    """Envoltorio de una lista de filas con la misma forma que
    `Worksheet.iter_rows(values_only=True)` de openpyxl."""

    def __init__(self, filas: list[list[str]]):
        self._filas = filas

    def iter_rows(self, min_row: int = 1, max_row: int | None = None,
                 values_only: bool = True):
        fin = len(self._filas) if max_row is None else max_row
        for fila in self._filas[min_row - 1:fin]:
            yield tuple(fila)


class WorkbookCSV:
    """Envoltorio mínimo para que un .csv se pueda leer con la misma API
    que un `Workbook` de openpyxl (`sheetnames` + `__getitem__().iter_rows`)
    — así el resto de este módulo (sheet_names, read_headers,
    guess_mapping_generic, read_rows_generic) no necesita duplicarse ni
    saber si el origen es Excel o CSV."""

    def __init__(self, filas: list[list[str]]):
        self.sheetnames = ["CSV"]
        self._hoja = _HojaCSV(filas)

    def __getitem__(self, _nombre_hoja):
        return self._hoja

    def close(self) -> None:
        pass  # nada que cerrar: ya se leyó todo el archivo en memoria


def es_csv(path: str) -> bool:
    return path.lower().endswith(".csv")


def abrir_csv(path: str, delimitador: str = ",") -> WorkbookCSV:
    with open(path, encoding="utf-8-sig", newline="") as f:
        filas = [list(fila) for fila in csv.reader(f, delimiter=delimitador)]
    return WorkbookCSV(filas)


def abrir_archivo_importacion(path: str):
    """Punto de entrada único para 'Importar cambios': abre .xlsx/.xlsm con
    openpyxl o .csv con abrir_csv(), devolviendo en ambos casos algo
    compatible con sheet_names/read_headers/read_rows_generic."""
    if es_csv(path):
        return abrir_csv(path)
    return open_workbook(path)


def sheet_names(workbook) -> list[str]:
    return list(workbook.sheetnames)


def read_headers(workbook, sheet: str) -> list[str]:
    """Devuelve los encabezados (fila 1) de la hoja indicada."""
    ws = workbook[sheet]
    row_iter = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    first_row = next(row_iter, ())
    headers = []
    for i, h in enumerate(first_row):
        text = str(h).strip() if h is not None else ""
        headers.append(text if text else f"(columna {i + 1})")
    return headers


def normalize_code(raw: object) -> str:
    """Convierte el valor crudo de una celda de código a texto comparable.
    Excel suele guardar códigos numéricos como número (p. ej. 4981005962.0),
    perdiendo los ceros a la izquierda; acá se normaliza a texto entero."""
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    return str(raw).strip()


def _normalizar_separador_decimal(text: str) -> str:
    """Resuelve el caso ambiguo de un número con AMBOS separadores
    presentes (p. ej. "1.234,56" en formato regional, o "1,234.56" en
    formato US): el que aparece último es el separador decimal, el otro es
    de miles y se descarta. Si solo aparece uno de los dos, se lo trata
    como separador decimal (comportamiento previo, sin cambios) — distinguir
    "1.234" (¿mil doscientos treinta y cuatro, o 1,234?) sin más contexto
    es ambiguo y no se intenta adivinar."""
    tiene_punto = "." in text
    tiene_coma = "," in text
    if tiene_punto and tiene_coma:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
        return text
    if tiene_coma:
        return text.replace(",", ".")
    return text


def tuvo_decimales(raw: object) -> bool:
    """True si `raw` (antes de redondear) representaba un número con parte
    fraccionaria no nula — para poder avisar cuando se lo importa a un
    campo entero y el redondeo silencioso (`int(round(...))`) le hace
    perder información (Nivel 4.6: 'aviso de redondeo')."""
    if raw is None:
        return False
    if isinstance(raw, bool):
        return False
    if isinstance(raw, int):
        return False
    if isinstance(raw, float):
        return not raw.is_integer()
    text = _normalizar_separador_decimal(str(raw).strip())
    if not text:
        return False
    try:
        return not float(text).is_integer()
    except ValueError:
        return False


def normalize_int(raw: object) -> int | None:
    """Convierte el valor crudo de una celda numérica (Color, Gramos,
    Velocidad) a int. Devuelve None si la celda está vacía o no es
    interpretable, para no pisar el valor actual con basura."""
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return int(round(raw))
    text = _normalizar_separador_decimal(str(raw).strip())
    if not text:
        return None
    try:
        return int(round(float(text)))
    except ValueError:
        return None


def normalize_decimal(raw: object) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    text = _normalizar_separador_decimal(str(raw).strip())
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def normalize_value(raw: object, tipo: str):
    """Normaliza una celda según el tipo del campo del perfil. Los campos de
    texto conservan ceros a la izquierda; los numéricos devuelven None si la
    celda no aporta un valor (para no pisar el dato actual)."""
    if tipo == "texto":
        return normalize_code(raw)
    if tipo == "decimal":
        return normalize_decimal(raw)
    return normalize_int(raw)


# --- Mapeo genérico guiado por los campos del perfil -------------------------
def guess_mapping_generic(headers: list[str], campos) -> dict[str, int | None]:
    """Para cada campo del perfil, sugiere el índice de columna del Excel que
    mejor coincide con su título/etiqueta/nombre. Usa coincidencia por
    subcadena y, si no, similitud aproximada (difflib, stdlib). Asigna de
    forma golosa por mejor puntaje y sin repetir columnas."""
    norm_headers = [_strip_accents(h.lower()) for h in headers]

    scored: list[tuple[float, str, int]] = []
    for campo in campos:
        terms = set()
        for t in (campo.titulo_ui, campo.nombre_interno, campo.etiqueta):
            for w in _strip_accents(str(t).lower()).replace("_", " ").split():
                if len(w) >= 3:
                    terms.add(w)
        for hi, h in enumerate(norm_headers):
            best = 0.0
            for term in terms:
                if term and (term in h or h in term):
                    best = max(best, 0.92)
                else:
                    best = max(best, difflib.SequenceMatcher(None, term, h).ratio())
            scored.append((best, campo.nombre_interno, hi))

    scored.sort(reverse=True)
    result: dict[str, int | None] = {}
    usadas: set[int] = set()
    for score, field, hi in scored:
        if field in result or hi in usadas:
            continue
        if score >= 0.55:
            result[field] = hi
            usadas.add(hi)
    for campo in campos:
        result.setdefault(campo.nombre_interno, None)
    return result


def read_rows_generic(workbook, sheet: str,
                      col_by_field: dict[str, int]) -> list[dict]:
    """Lee las filas de datos devolviendo, por fila, un dict
    {nombre_interno: valor_crudo} según el mapeo de columnas confirmado."""
    ws = workbook[sheet]
    out: list[dict] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if raw is None or all(c is None for c in raw):
            continue
        rec = {}
        for field, ci in col_by_field.items():
            rec[field] = raw[ci] if ci is not None and ci < len(raw) else None
        out.append(rec)
    return out
