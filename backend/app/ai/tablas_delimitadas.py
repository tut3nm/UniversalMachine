"""
Tabula archivos de texto delimitados que traen varias tablas en el mismo
archivo, separadas por lineas de titulo ("[Messprogrammseite 1]", "# QSStat")
o lineas vacias: el formato de los CSV del equipo H1312 (docs/H1312/).

Determinista y sin IA. El asistente puede corregir la estructura detectada
(rangos de filas, separador) via el DSL, pero lo que termina en el Excel sale
siempre del archivo. Ver PLAN_ASISTENTE_IA.md, seccion 12.1.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SEPARADORES = (",", ";", "\t", "|")

_LINEAS_MUESTRA_SEPARADOR = 500
_EXCEL_MAX_DIGITOS = 15
_ENTERO_RE = re.compile(r"^-?\d+$")
_DECIMAL_RE = re.compile(r"^-?\d+\.\d+$")
_CARACTERES_INVALIDOS_HOJA = re.compile(r"[\[\]:*?/\\]")


@dataclass(frozen=True)
class TablaDef:
    """Rango de lineas (1-based, inclusivo) de una tabla del archivo."""
    desde: int
    hasta: int
    con_encabezado: bool = True


@dataclass
class TablaLeida:
    titulo: str
    desde: int
    hasta: int
    con_encabezado: bool
    fila_encabezado: int | None
    columnas: list[str]
    filas: list[list] = field(default_factory=list)


@dataclass
class ResultadoTablas:
    separador: str
    tablas: list[TablaLeida]
    advertencias: list[str] = field(default_factory=list)


def decodificar(contenido: bytes) -> list[str]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = contenido.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        texto = contenido.decode("latin-1", errors="replace")
    return texto.splitlines()


def es_titulo(linea: str) -> bool:
    s = linea.strip()
    return (s.startswith("[") and s.endswith("]") and len(s) > 2) or s.startswith("#")


def _texto_titulo(linea: str) -> str:
    s = linea.strip()
    if s.startswith("[") and s.endswith("]"):
        return s[1:-1].strip()
    return s.lstrip("#").strip()


def _es_contenido(linea: str) -> bool:
    return linea.strip() != "" and not es_titulo(linea)


def detectar_separador(lineas: list[str]) -> str | None:
    """El separador de la lista cerrada que aparece en mas lineas de
    contenido. None si ninguno aparece en al menos la mitad: no es un archivo
    delimitado (ej. los logs Prod.*.TXT, que no tienen ninguno)."""
    muestra = [l for l in lineas[:_LINEAS_MUESTRA_SEPARADOR] if _es_contenido(l)]
    if not muestra:
        return None
    apariciones = {sep: sum(1 for l in muestra if sep in l) for sep in SEPARADORES}
    mejor = max(SEPARADORES, key=lambda s: apariciones[s])
    if apariciones[mejor] < 2 or apariciones[mejor] * 2 < len(muestra):
        return None
    return mejor


def _campos(linea: str, sep: str) -> list[str]:
    return [c.strip() for c in next(csv.reader([linea], delimiter=sep))]


def _es_numero(texto: str) -> bool:
    return bool(_ENTERO_RE.match(texto) or _DECIMAL_RE.match(texto))


def _es_clave_valor(filas: list[list[str]]) -> bool:
    return all(len(f) <= 2 and f and f[0] and not _es_numero(f[0]) for f in filas)


def detectar_tablas(lineas: list[str], sep: str) -> list[TablaDef]:
    """Un bloque por cada tramo continuo de lineas de contenido (los cortan
    las lineas vacias y los titulos). Un bloque donde todas las filas son
    'clave,valor' no tiene encabezado; el resto usa su primera fila."""
    bloques: list[tuple[int, int]] = []
    inicio: int | None = None
    for n, linea in enumerate(lineas, start=1):
        if _es_contenido(linea):
            if inicio is None:
                inicio = n
        elif inicio is not None:
            bloques.append((inicio, n - 1))
            inicio = None
    if inicio is not None:
        bloques.append((inicio, len(lineas)))

    defs: list[TablaDef] = []
    for desde, hasta in bloques:
        filas = [_campos(lineas[n - 1], sep) for n in range(desde, hasta + 1)]
        if _es_clave_valor(filas):
            defs.append(TablaDef(desde, hasta, con_encabezado=False))
        else:
            con_encabezado = len(filas) >= 2 and any(c and not _es_numero(c) for c in filas[0])
            defs.append(TablaDef(desde, hasta, con_encabezado=con_encabezado))
    return defs


def es_bloque_clave_valor(lineas: list[str], sep: str, desde: int, hasta: int) -> bool:
    """Respaldo determinista de la clasificacion listado/tabla: True si todas
    las filas de contenido del rango tienen forma clave,valor. Se usa cuando
    el modelo no clasifico el rango (no respondio, o no incluyo 'tipo')."""
    filas = [_campos(lineas[n - 1], sep) for n in range(desde, hasta + 1) if _es_contenido(lineas[n - 1])]
    return bool(filas) and _es_clave_valor(filas)


def es_formato_tablas(lineas: list[str]) -> bool:
    """True si el archivo tiene la forma que este motor sabe tabular: un
    separador reconocible y al menos una tabla con encabezado y datos."""
    sep = detectar_separador(lineas)
    if sep is None:
        return False
    return any(d.con_encabezado and d.hasta > d.desde for d in detectar_tablas(lineas, sep))


def _convertir(valor: str):
    s = valor.strip()
    if s == "":
        return None
    digitos = s.lstrip("-")
    con_cero_adelante = len(digitos) > 1 and digitos[0] == "0" and digitos[1] != "."
    if con_cero_adelante or len(digitos.replace(".", "")) > _EXCEL_MAX_DIGITOS:
        return s
    if _ENTERO_RE.match(s):
        return int(s)
    if _DECIMAL_RE.match(s):
        return float(s)
    return s


def validar_defs(defs: list[TablaDef], total_lineas: int) -> None:
    if not defs:
        raise ValueError("No hay ninguna tabla definida.")
    for d in defs:
        if not (1 <= d.desde <= d.hasta <= total_lineas):
            raise ValueError(
                f"El rango de filas {d.desde}-{d.hasta} no es valido: el archivo tiene "
                f"{total_lineas} filas."
            )
    ordenadas = sorted(defs, key=lambda d: d.desde)
    for a, b in zip(ordenadas, ordenadas[1:]):
        if b.desde <= a.hasta:
            raise ValueError(
                f"Las tablas de las filas {a.desde}-{a.hasta} y {b.desde}-{b.hasta} se superponen."
            )


def texto_titulo(linea: str) -> str:
    return _texto_titulo(linea)


def titulo_para(lineas: list[str], desde: int, hasta: int) -> str | None:
    """Titulo de seccion del bloque desde-hasta (el de la propia primera
    linea, o el mas cercano hacia arriba sin contenido en el medio). Usado
    tambien por app/ai/memoria/anclaje.py para anclar/resolver reglas."""
    return _titulo_para(lineas, TablaDef(desde, hasta))


def _titulo_para(lineas: list[str], d: TablaDef) -> str | None:
    for n in range(d.desde, d.hasta + 1):
        linea = lineas[n - 1]
        if es_titulo(linea):
            return _texto_titulo(linea)
        if linea.strip():
            break
    n = d.desde - 1
    while n >= 1:
        linea = lineas[n - 1]
        if es_titulo(linea):
            return _texto_titulo(linea)
        if linea.strip():
            return None
        n -= 1
    return None


def leer_tablas(lineas: list[str], sep: str, defs: list[TablaDef]) -> ResultadoTablas:
    if sep not in SEPARADORES:
        raise ValueError(f"Separador no soportado: {sep!r}.")
    validar_defs(defs, len(lineas))

    advertencias: list[str] = []
    tablas: list[TablaLeida] = []
    for i, d in enumerate(sorted(defs, key=lambda d: d.desde), start=1):
        titulo = _titulo_para(lineas, d) or f"Tabla {i}"
        contenido: list[tuple[int, list[str]]] = []
        for n in range(d.desde, d.hasta + 1):
            linea = lineas[n - 1]
            if es_titulo(linea):
                if contenido:
                    advertencias.append(
                        f"{titulo}: la fila {n} ('{linea.strip()}') es un titulo de seccion; "
                        "no se tomo como dato."
                    )
                continue
            if linea.strip():
                contenido.append((n, _campos(linea, sep)))
        if not contenido:
            raise ValueError(f"La tabla de las filas {d.desde}-{d.hasta} no tiene datos.")

        if d.con_encabezado:
            fila_encabezado, encabezado = contenido[0]
            datos = [c for _, c in contenido[1:]]
            if not datos:
                raise ValueError(
                    f"La tabla de las filas {d.desde}-{d.hasta} solo tiene la fila de encabezado."
                )
            if fila_encabezado != d.desde:
                advertencias.append(f"{titulo}: se uso la fila {fila_encabezado} como encabezado.")
        else:
            fila_encabezado, encabezado = None, []
            datos = [c for _, c in contenido]

        n_columnas = max(len(encabezado), max(len(f) for f in datos))
        while (
            n_columnas > 0
            and (n_columnas > len(encabezado) or not encabezado[n_columnas - 1])
            and all(len(f) < n_columnas or f[n_columnas - 1] == "" for f in datos)
        ):
            n_columnas -= 1

        if not d.con_encabezado and n_columnas == 2:
            columnas = ["Campo", "Valor"]
        else:
            columnas = [
                encabezado[j] if j < len(encabezado) and encabezado[j] else f"Columna {j + 1}"
                for j in range(n_columnas)
            ]
            sin_nombre = [c for j, c in enumerate(columnas) if j >= len(encabezado) or not encabezado[j]]
            if d.con_encabezado and sin_nombre:
                advertencias.append(
                    f"{titulo}: el encabezado tiene {len(encabezado)} nombres pero hay "
                    f"{n_columnas} columnas; las que faltan se llamaron "
                    f"{sin_nombre[0]}...{sin_nombre[-1]}."
                )

        filas = [
            [_convertir(f[j]) if j < len(f) else None for j in range(n_columnas)]
            for f in datos
        ]
        tablas.append(TablaLeida(
            titulo=titulo, desde=d.desde, hasta=d.hasta, con_encabezado=d.con_encabezado,
            fila_encabezado=fila_encabezado, columnas=columnas, filas=filas,
        ))
    return ResultadoTablas(separador=sep, tablas=tablas, advertencias=advertencias)


def _nombre_hoja(titulo: str, usados: set[str]) -> str:
    base = _CARACTERES_INVALIDOS_HOJA.sub(" ", titulo).strip()[:31] or "Tabla"
    nombre, n = base, 2
    while nombre.lower() in usados:
        sufijo = f" ({n})"
        nombre = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(nombre.lower())
    return nombre


def escribir_excel(resultado: ResultadoTablas) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    usados: set[str] = set()
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="DDE6F0")
    for tabla in resultado.tablas:
        ws = wb.create_sheet(_nombre_hoja(tabla.titulo, usados))
        ws.append(tabla.columnas)
        for celda in ws[1]:
            celda.font = negrita
            celda.fill = relleno
        for fila in tabla.filas:
            ws.append(fila)
        ws.freeze_panes = "A2"
        if tabla.filas:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(tabla.columnas))}{len(tabla.filas) + 1}"
        for j, nombre in enumerate(tabla.columnas, start=1):
            ancho = max([len(str(nombre))] + [len(str(f[j - 1])) for f in tabla.filas[:200] if f[j - 1] is not None])
            ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 2, 8), 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def resumen_tabla(tabla: TablaLeida) -> dict:
    return {
        "titulo": tabla.titulo,
        "desde": tabla.desde,
        "hasta": tabla.hasta,
        "con_encabezado": tabla.con_encabezado,
        "fila_encabezado": tabla.fila_encabezado,
        "n_filas": len(tabla.filas),
        "n_columnas": len(tabla.columnas),
        "columnas": tabla.columnas[:8],
    }
