"""
Exportador a Excel generico: funciona con CUALQUIER formato que structure.py
haya podido descubrir, usando las etiquetas que puso labeler.py.

Genera tres hojas:
  Resumen    -> datos generales y los parametros de configuracion encontrados
  Mediciones -> una fila por registro, con los valores tal cual los dio el equipo
  Estructura -> que fue lo que el programa detecto (para poder auditarlo)
"""
from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1F4E78")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E78")
VALUE_FONT = Font(name="Calibri", size=10)
NOTE_FONT = Font(name="Calibri", size=9, italic=True, color="808080")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_FONT = Font(name="Calibri", size=10)
TOTALS_FILL = PatternFill("solid", fgColor="FFF2CC")
AI_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _num(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _header_text(label_obj) -> str:
    if label_obj.unit:
        return f"{label_obj.label} ({label_obj.unit})"
    return label_obj.label


def _write_resumen(wb, discovery, labels, source_path, annotated_path):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False

    row = 1
    ws.cell(row=row, column=1, value="Datos de ensayo exportados a Excel").font = TITLE_FONT
    row += 1
    ws.cell(row=row, column=1, value=f"Archivo de datos: {os.path.basename(source_path)}").font = SUBTITLE_FONT
    row += 1
    if annotated_path:
        ws.cell(row=row, column=1,
                value=f"Archivo con anotaciones: {os.path.basename(annotated_path)}").font = SUBTITLE_FONT
        row += 1
    ws.cell(row=row, column=1,
            value=f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}").font = SUBTITLE_FONT
    row += 2

    ws.cell(row=row, column=1, value="RESUMEN").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    for c in range(2, 7):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1

    fechas = [r.fields.get("timestamp") for r in discovery.records
              if isinstance(r.fields.get("timestamp"), datetime)]
    pares = [
        ("Mediciones encontradas", len(discovery.records)),
        ("Campos por medicion", len(discovery.record_field_names)),
        ("Bloques de configuracion", len(discovery.config_snapshots)),
        ("Lineas del archivo", discovery.total_lines),
        ("Lineas no interpretadas", len(discovery.unparsed_lines)),
    ]
    if fechas:
        pares.append(("Primera medicion", fechas[0].strftime("%d/%m/%Y %H:%M:%S")))
        pares.append(("Ultima medicion", fechas[-1].strftime("%d/%m/%Y %H:%M:%S")))
    for k, v in pares:
        ws.cell(row=row, column=1, value=k).font = LABEL_FONT
        c = ws.cell(row=row, column=2, value=v)
        c.font = VALUE_FONT
        c.fill = TOTALS_FILL
        row += 1

    row += 1
    for msg in labels.messages:
        ws.cell(row=row, column=1, value=msg).font = NOTE_FONT
        row += 1
    ws.cell(row=row, column=1, value=(
        "Los nombres en verde fueron propuestos por la IA local a partir del archivo "
        "con anotaciones. Los numeros NO pasan por la IA: se leen de forma directa "
        "del archivo original."
    )).font = NOTE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 28
    row += 2

    # ---- configuraciones -------------------------------------------------
    ws.cell(row=row, column=1, value="PARAMETROS DE CONFIGURACION DEL ENSAYO").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    for c in range(2, 10):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1

    keys = discovery.config_keys
    headers = ["Bloque", "Desde la linea", "Mediciones"] + [
        _header_text(labels.config_label(k)) for k in keys
    ]
    hrow = row
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=hrow, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
        if c > 3 and labels.config_label(keys[c - 4]).source == "ia":
            cell.fill = PatternFill("solid", fgColor="375623")
    ws.row_dimensions[hrow].height = 32
    row += 1

    for snap in discovery.config_snapshots:
        vals = [snap.index, snap.line_index + 1, snap.n_records]
        for k in keys:
            e = snap.entries.get(k)
            if e is None:
                vals.append("")
            elif e.numbers:
                vals.append(", ".join(str(_num(v)) for v in e.numbers))
            else:
                vals.append(e.text)
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = DATA_FONT
            cell.border = BORDER
        row += 1

    # ---- claves originales ----------------------------------------------
    row += 1
    ws.cell(row=row, column=1, value="Clave original en el archivo:").font = LABEL_FONT
    for c, k in enumerate(keys, start=4):
        cell = ws.cell(row=row, column=c, value=k)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(horizontal="center")

    widths = {"A": 26, "B": 15, "C": 12}
    for i in range(len(keys)):
        widths[get_column_letter(4 + i)] = 22
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _write_mediciones(wb, discovery, labels, progress_callback=None):
    ws = wb.create_sheet("Mediciones")
    names = discovery.record_field_names

    headers = ["N°"] + [_header_text(labels.field_label(n)) for n in names]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = BORDER
        if c > 1 and labels.field_label(names[c - 2]).source == "ia":
            cell.fill = PatternFill("solid", fgColor="375623")
    ws.row_dimensions[1].height = 34

    # segunda fila: el nombre crudo que usa el equipo
    ws.cell(row=2, column=1, value="").font = NOTE_FONT
    for c, n in enumerate(names, start=2):
        cell = ws.cell(row=2, column=c, value=n)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    n_total = len(discovery.records)
    report_every = max(1, n_total // 100)

    for i, rec in enumerate(discovery.records):
        if progress_callback and i % report_every == 0:
            progress_callback(i, n_total)
        r = i + 3
        ws.cell(row=r, column=1, value=rec.number)
        for c, name in enumerate(names, start=2):
            v = rec.fields.get(name)
            if isinstance(v, datetime):
                cell = ws.cell(row=r, column=c, value=v)
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            else:
                ws.cell(row=r, column=c, value=_num(v))

    widths = {"A": 8}
    for i in range(len(names)):
        widths[get_column_letter(2 + i)] = 18
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    if progress_callback:
        progress_callback(n_total, n_total)


def _write_estructura(wb, discovery, labels):
    ws = wb.create_sheet("Estructura detectada")
    ws.sheet_view.showGridLines = False
    row = 1
    ws.cell(row=row, column=1, value="Que detecto el programa en este archivo").font = TITLE_FONT
    row += 2
    ws.cell(row=row, column=1, value=(
        "Esta hoja sirve para auditar el resultado: muestra como se interpreto el "
        "formato. Si algo esta mal, conviene revisar aca antes que en los datos."
    )).font = NOTE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 30
    row += 2

    datos = [
        ("Lineas totales", discovery.total_lines),
        ("Lineas de datos", discovery.line_kind_counts.get("data", 0)),
        ("Lineas de configuracion", discovery.line_kind_counts.get("meta", 0)),
        ("Lineas separadoras", discovery.line_kind_counts.get("delim", 0)),
        ("Lineas en blanco", discovery.line_kind_counts.get("blank", 0)),
        ("Lineas por medicion (ciclo detectado)", discovery.period),
        ("Confianza del ciclo", f"{discovery.period_score*100:.1f}%"),
        ("Patron que inicia cada medicion", discovery.record_start_signature or "-"),
    ]
    for k, v in datos:
        ws.cell(row=row, column=1, value=k).font = LABEL_FONT
        ws.cell(row=row, column=2, value=v).font = VALUE_FONT
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="CAMPOS DE CADA MEDICION").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    for c in range(2, 5):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1
    for c, h in enumerate(["Clave del equipo", "Nombre asignado", "Unidad", "Origen del nombre"], start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
    row += 1
    for n in discovery.record_field_names:
        lb = labels.field_label(n)
        vals = [n, lb.label, lb.unit, "IA local" if lb.source == "ia" else "nombre original"]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=c, value=v)
            cell.font = DATA_FONT
            cell.border = BORDER
            if lb.source == "ia":
                cell.fill = AI_FILL
        row += 1

    if discovery.unparsed_lines:
        row += 1
        ws.cell(row=row, column=1, value="LINEAS QUE NO SE PUDIERON INTERPRETAR").font = SECTION_FONT
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor="C00000")
        for c in range(2, 5):
            ws.cell(row=row, column=c).fill = PatternFill("solid", fgColor="C00000")
        row += 1
        for idx, raw in discovery.unparsed_lines:
            ws.cell(row=row, column=1, value=idx + 1).font = DATA_FONT
            ws.cell(row=row, column=2, value=raw.strip()[:200]).font = DATA_FONT
            row += 1

    for col, w in {"A": 34, "B": 34, "C": 14, "D": 18}.items():
        ws.column_dimensions[col].width = w


def write_excel(discovery, labels, output_path, source_path,
                annotated_path=None, progress_callback=None):
    wb = Workbook()
    _write_resumen(wb, discovery, labels, source_path, annotated_path)
    _write_mediciones(wb, discovery, labels, progress_callback=progress_callback)
    _write_estructura(wb, discovery, labels)
    wb.active = 0
    wb.save(output_path)
