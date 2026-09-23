"""
Tests de app/ai/csv_recipe.py (migrado de Diagramadora/ai/csv_recipe.py,
PLAN_EDITOR_RECETAS_MATRIZ.md fase A). No habia tests previos: se escriben
desde cero, con foco en el contrato de seguridad del modulo (solo cambian
las celdas editadas, todo lo demas queda byte a byte igual) y en el camino
de ALERTA (autoverificacion post-escritura), que es lo que en la fase B va a
bloquear la descarga.
"""
import os

import pytest

from app.ai import csv_recipe as R

# 6 productos (no 3): _looks_like_index_row exige al menos 5 valores para
# activarse, y el formato real siempre trae esa fila de indices "3,1,2,..."
# despues de la fila de codigos de producto. 3 filas de parametro (no 2):
# sniff() exige que al menos 70% de las lineas tengan la cantidad de
# columnas "comun" (aca 7) para reconocer la matriz, y con solo 2 no llega.
# Los decimales evitan terminar en ".0": write_excel_recipe() guarda ese
# valor como ENTERO en la celda (_as_excel_value), asi que "3.0" volveria de
# Excel como "3" aunque nadie lo haya tocado - no es un bug de la migracion,
# es como ya se comportaba csv_recipe.py.
CSV_TEXTO = (
    "List separator=,Decimal symbol=.\r\n"
    "Recipe_1\r\n"
    "LANGID_409,P1,P2,P3,P4,P5,P6\r\n"
    "3,1,2,3,4,5,6\r\n"
    "Temperatura,45,50,55,60,65,70\r\n"
    "Presion,1.5,2.2,2.7,3.1,3.6,4.4\r\n"
    "Dureza,10,20,30,40,50,60\r\n"
)


def _escribir(tmp_path, nombre="receta.csv", texto=CSV_TEXTO):
    p = tmp_path / nombre
    p.write_bytes(texto.encode("utf-8"))
    return str(p)


# --- sniff / read_recipe -----------------------------------------------


def test_sniff_detecta_matriz_de_receta(tmp_path):
    path = _escribir(tmp_path)
    ok, msg = R.sniff(path)
    assert ok
    assert "7" in msg  # 7 columnas: LANGID + 6 productos


def test_sniff_rechaza_archivo_sin_forma_de_matriz(tmp_path):
    path = tmp_path / "no_es_matriz.csv"
    path.write_text("una,linea\notra\n")
    ok, _ = R.sniff(str(path))
    assert not ok


def test_read_recipe_detecta_productos_y_parametros(tmp_path):
    rec = R.read_recipe(_escribir(tmp_path))
    assert rec.product_codes == ["P1", "P2", "P3", "P4", "P5", "P6"]
    assert [p.name for p in rec.params] == ["Temperatura", "Presion", "Dureza"]
    assert rec.line_ending == "\r\n"


def test_read_recipe_trata_fila_de_indices_como_extra_no_como_parametro(tmp_path):
    rec = R.read_recipe(_escribir(tmp_path))
    nombres = [p.name for p in rec.params]
    assert "3" not in nombres
    assert len(rec.extra_rows) == 1


def test_read_recipe_sin_fila_de_productos_lanza_error(tmp_path):
    path = tmp_path / "sin_productos.csv"
    path.write_text("a,b\nc,d\n")
    with pytest.raises(ValueError):
        R.read_recipe(str(path))


# --- rebuild_csv: contrato de seguridad ---------------------------------


def test_rebuild_csv_sin_ediciones_es_byte_identico(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")

    res = R.rebuild_csv(original, rec, {}, salida)

    assert res.ok
    assert res.changes == []
    with open(original, "rb") as fo, open(salida, "rb") as fs:
        assert fo.read() == fs.read()


def test_rebuild_csv_aplica_una_edicion_y_deja_el_resto_igual(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")

    res = R.rebuild_csv(original, rec, {("Temperatura", "P2"): "99"}, salida)

    assert res.ok
    assert len(res.changes) == 1
    cambio = res.changes[0]
    assert (cambio.parametro, cambio.producto, cambio.valor_anterior, cambio.valor_nuevo) == (
        "Temperatura", "P2", "50", "99",
    )

    verif = R.read_recipe(salida)
    valores_temp = dict(zip(verif.product_codes, next(p.values for p in verif.params if p.name == "Temperatura")))
    assert valores_temp["P2"] == "99"
    valores_pres = dict(zip(verif.product_codes, next(p.values for p in verif.params if p.name == "Presion")))
    assert valores_pres == {"P1": "1.5", "P2": "2.2", "P3": "2.7", "P4": "3.1", "P5": "3.6", "P6": "4.4"}


def test_rebuild_csv_preserva_crlf_original(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")
    R.rebuild_csv(original, rec, {("Temperatura", "P1"): "1"}, salida)
    with open(salida, "rb") as f:
        assert b"\r\n" in f.read()


def test_rebuild_csv_codigo_de_producto_desconocido_no_se_aplica_y_avisa(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")

    res = R.rebuild_csv(original, rec, {("Temperatura", "P999"): "1"}, salida)

    assert res.changes == []
    assert any("P999" in w for w in res.warnings)
    assert res.ok  # es un aviso, no una ALERTA: el resto del archivo sigue intacto


def test_rebuild_csv_valor_vacio_avisa(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")

    res = R.rebuild_csv(original, rec, {("Temperatura", "P1"): ""}, salida)

    assert any("valor vacío" in w for w in res.warnings)


def test_rebuild_csv_texto_donde_siempre_hubo_numero_avisa(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")

    res = R.rebuild_csv(original, rec, {("Temperatura", "P1"): "ROTO"}, salida)

    assert any("no es numérico" in w for w in res.warnings)
    # igual se aplica (es un aviso a revisar, no un bloqueo del motor)
    assert res.changes and res.changes[0].valor_nuevo == "ROTO"


def test_rebuild_csv_no_pisa_filas_de_parametros_no_editados(tmp_path):
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")
    R.rebuild_csv(original, rec, {("Temperatura", "P1"): "111"}, salida)

    with open(original, encoding="utf-8") as f:
        lineas_originales = f.read().splitlines()
    with open(salida, encoding="utf-8") as f:
        lineas_salida = f.read().splitlines()

    # todas las lineas excepto la de "Temperatura" quedan exactamente iguales
    idx_temp = next(i for i, l in enumerate(lineas_originales) if l.startswith("Temperatura"))
    for i in range(len(lineas_originales)):
        if i == idx_temp:
            continue
        assert lineas_originales[i] == lineas_salida[i]


# --- ALERTA: la autoverificacion post-escritura -------------------------


def test_rebuild_csv_valor_con_salto_de_linea_embebido_dispara_alerta(tmp_path):
    """Un valor editado con un '\\n' real adentro corre las lineas fisicas
    del archivo de salida: la autoverificacion lo tiene que cazar. Esto es
    justo lo que en la fase B bloquea la descarga (PLAN_EDITOR_RECETAS_MATRIZ.md
    seccion 2, decision 2)."""
    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    salida = str(tmp_path / "salida.csv")

    res = R.rebuild_csv(original, rec, {("Temperatura", "P1"): "5\n6"}, salida)

    assert not res.ok
    assert any(w.startswith("ALERTA") for w in res.warnings)


# --- Excel: round-trip completo (export -> editar -> reimportar) --------


def test_round_trip_completo_via_excel(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")

    original = _escribir(tmp_path)
    rec = R.read_recipe(original)
    xlsx_path = str(tmp_path / "receta.xlsx")
    R.write_excel_recipe(rec, xlsx_path, original)
    assert os.path.exists(xlsx_path)

    # simular la edicion que haria un operario en Excel: cambiar una celda
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[R.SHEET_RECETAS]
    assert ws.cell(row=1, column=1).value == R.PRODUCT_COL
    # fila 2 = P1 (primer producto), buscar la columna "Temperatura"
    headers = [c.value for c in ws[1]]
    col_temp = headers.index("Temperatura") + 1
    ws.cell(row=2, column=col_temp, value=123)
    wb.save(xlsx_path)

    editado = R.read_edited_excel(xlsx_path)
    assert editado[("Temperatura", "P1")] == "123"

    salida = str(tmp_path / "salida.csv")
    res = R.rebuild_csv(original, rec, editado, salida)

    assert res.ok
    assert len(res.changes) == 1
    assert res.changes[0] == R.CellChange("Temperatura", "P1", "45", "123")


def test_read_edited_excel_sin_hoja_recetas_lanza_error(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    xlsx_path = str(tmp_path / "vacio.xlsx")
    Workbook().save(xlsx_path)
    with pytest.raises(ValueError):
        R.read_edited_excel(xlsx_path)
