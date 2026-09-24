import io
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app.ai import llm
from app.services import asistente_service as svc

REPO_ROOT = Path(__file__).resolve().parents[2]
LISTADO_PATH = REPO_ROOT / "docs" / "Recetas_Andon.txt"
PLANTILLA_PATH = REPO_ROOT / "docs" / "Recetas" / "plantilla.txt"

# El pedido real del 2026-09-23 en /mediciones (PLAN_ASISTENTE_IA.md, seccion 12).
PEDIDO_REAL = (
    "Tabulá las mediciones de este ensayo a un Excel. Vas a ver que de la fila 1 a la 12 "
    "corresponde a la primera tabla de medición, donde cada fila corresponde a un registro, y "
    "la primer fila aclara el nombre de las columnas. De la fila 14 a 212 es la otra tabla de "
    "registros, donde la primera fila es el nombre de la columna y el resto de las filas es "
    "cada uno de los registros de las tablas. Las columnas están separadas por \",\" en ambas tablas."
)


def _csv_como_el_del_pedido() -> bytes:
    """Misma forma y numeracion que el archivo que subio el usuario: titulo
    (1), encabezado (2), 9 filas (3-11), titulo (12), vacia (13), encabezado
    (14) y 198 filas (15-212) con coma final."""
    lineas = ["[Messprogrammseite 1]", "Zyklen,Hub,Mittelpos,Geschw,ZugOG-A1,ZugUG-A1,DruckOG-A1,DruckUG-A1"]
    lineas += [f"{i},75,1335,0.262,0,0,0,0,4621242,4613528,4200380,4580856" for i in range(9)]
    lineas += ["# QSStat", "", "Datum,Uhrzeit,Achse,Bewertung,Nummer,ZugV1,DruckV1"]
    lineas += [f"28.08.26,04:37:28,A1,IO,{i},1994,   0," for i in range(1, 199)]
    assert len(lineas) == 212
    return "\r\n".join(lineas).encode("utf-8")


def _archivos_mediciones() -> dict[str, svc.Archivo]:
    return {"datos": svc.Archivo("824902015333_280826_006.csv", _csv_como_el_del_pedido())}


def _archivos_recetas() -> dict[str, svc.Archivo]:
    return {"listado": svc.Archivo(LISTADO_PATH.name, LISTADO_PATH.read_bytes())}


def _archivos_plantilla() -> dict[str, svc.Archivo]:
    return {
        "plantilla": svc.Archivo("plantilla.txt", PLANTILLA_PATH.read_bytes()),
        "listado": svc.Archivo(LISTADO_PATH.name, LISTADO_PATH.read_bytes()),
    }


def _stub(respuesta: str):
    def ejecutar(system, user, grammar=None, max_tokens=768):
        return llm.LlmResult(True, respuesta)

    return ejecutar


def _no_deberia_llamarse(system, user, grammar=None, max_tokens=768):
    raise AssertionError("no hacia falta llamar a la IA")


def _llm_caido(system, user, grammar=None, max_tokens=768):
    return llm.LlmResult(False, "", "IA local no disponible.")


# -- generales ----------------------------------------------------------------


def test_pantalla_desconocida_se_rechaza():
    with pytest.raises(svc.AsistenteError, match="Pantalla desconocida"):
        svc.interpretar("otra", "hola", {}, ejecutar_llm=_no_deberia_llamarse)


def test_sin_los_archivos_de_la_pantalla_los_pide_sin_llamar_a_la_ia():
    r = svc.interpretar("mediciones", "tabula esto", {}, ejecutar_llm=_no_deberia_llamarse)
    assert r["faltan_archivos"] == ["el archivo de datos"]
    r = svc.interpretar("plantilla", "genera", _archivos_recetas(), ejecutar_llm=_no_deberia_llamarse)
    assert r["faltan_archivos"] == ["la plantilla"]


def test_ninguna_respuesta_deriva_a_otra_pantalla():
    r = svc.interpretar("mediciones", "cual es la capital de francia", _archivos_mediciones(),
                        ejecutar_llm=_stub("[]"))
    assert "requiere_pantalla" not in r
    assert r["sin_indicaciones"] is True
    assert len(r["tablas"]) == 2


# -- mediciones ---------------------------------------------------------------


def test_mediciones_pedido_real_da_las_dos_tablas_correctas():
    stub = _stub(
        '[{"op": "definir_tablas", "args": {"tablas": ['
        '{"desde": 1, "hasta": 12, "tipo": "tabla"},'
        ' {"desde": 14, "hasta": 212, "tipo": "tabla"}]}},'
        ' {"op": "usar_separador", "args": {"separador": ","}}]'
    )
    r = svc.interpretar("mediciones", PEDIDO_REAL, _archivos_mediciones(), ejecutar_llm=stub)
    assert r["ia_respondio"] and not r["sin_indicaciones"]
    assert [(t["titulo"], t["fila_encabezado"], t["n_filas"], t["n_columnas"]) for t in r["tablas"]] == [
        ("Messprogrammseite 1", 2, 9, 12),
        ("QSStat", 14, 198, 7),
    ]
    assert any("fila 12" in a for a in r["advertencias"])
    assert r["programa"][0] == {"op": "usar_separador", "args": {"separador": ","}}


def test_mediciones_rango_mencionado_y_no_usado_se_avisa():
    # El modelo solo tradujo la primera tabla que el usuario describio; la
    # segunda ("de la fila 14 a 212") se avisa como no usada.
    stub = _stub(
        '[{"op": "definir_tablas", "args": {"tablas": ['
        '{"desde": 1, "hasta": 12, "tipo": "tabla"}]}},'
        ' {"op": "usar_separador", "args": {"separador": ","}}]'
    )
    r = svc.interpretar("mediciones", PEDIDO_REAL, _archivos_mediciones(), ejecutar_llm=stub)
    assert any("14" in a and "212" in a for a in r["advertencias"])


def test_mediciones_sin_indicaciones_usa_la_deteccion_automatica():
    r = svc.interpretar("mediciones", "tabulalo", _archivos_mediciones(), ejecutar_llm=_stub("[]"))
    assert [(t["desde"], t["hasta"]) for t in r["tablas"]] == [(2, 11), (14, 212)]
    assert r["programa"][1]["op"] == "definir_tablas"


def test_mediciones_ia_caida_sigue_con_la_deteccion_y_lo_marca():
    r = svc.interpretar("mediciones", "tabulalo", _archivos_mediciones(), ejecutar_llm=_llm_caido)
    assert r["ia_respondio"] is False
    assert len(r["tablas"]) == 2


def test_mediciones_rango_fuera_del_archivo_vuelve_como_error_de_validacion():
    stub = _stub('[{"op": "definir_tablas", "args": {"tablas": [{"desde": 1, "hasta": 999, "tipo": "tabla"}]}}]')
    r = svc.interpretar("mediciones", "de la 1 a la 999", _archivos_mediciones(), ejecutar_llm=stub)
    assert "212 filas" in r["error_validacion"]


def test_mediciones_previsualizar_programa_editado():
    programa = [
        {"op": "usar_separador", "args": {"separador": ","}},
        {"op": "definir_tablas", "args": {"tablas": [{"desde": 14, "hasta": 50, "tipo": "tabla"}]}},
    ]
    r = svc.previsualizar("mediciones", programa, _archivos_mediciones())
    assert [(t["titulo"], t["n_filas"]) for t in r["tablas"]] == [("QSStat", 36)]


def test_mediciones_ejecutar_devuelve_excel_y_deja_log(tmp_path, monkeypatch):
    import paths
    from app.ai.dsl import log as dsl_log

    monkeypatch.setattr(paths, "app_base_dir", lambda: str(tmp_path))
    programa = [
        {"op": "usar_separador", "args": {"separador": ","}},
        {"op": "definir_tablas", "args": {"tablas": [
            {"desde": 1, "hasta": 12, "tipo": "tabla"},
            {"desde": 14, "hasta": 212, "tipo": "tabla"},
        ]}},
    ]
    salida = svc.ejecutar("mediciones", programa, _archivos_mediciones(), texto_usuario=PEDIDO_REAL)
    assert salida.nombre == "824902015333_280826_006.xlsx"
    assert salida.media_type == svc.MEDIA_XLSX
    wb = openpyxl.load_workbook(io.BytesIO(salida.contenido))
    assert wb.sheetnames == ["Messprogrammseite 1", "QSStat"]
    assert wb["QSStat"].max_row == 199

    ultimo = dsl_log.leer_eventos()[-1]
    assert ultimo["operacion"] == "mediciones"
    assert ultimo["programa"] == programa
    assert ultimo["texto_usuario"] == PEDIDO_REAL


def test_mediciones_operacion_de_otra_pantalla_se_rechaza():
    with pytest.raises(svc.AsistenteError, match="no es de mediciones"):
        svc.ejecutar("mediciones", [{"op": "filtrar_filas", "args": {}}], _archivos_mediciones())


# -- recetas por area -----------------------------------------------------------


def test_recetas_con_listado_devuelve_vista_previa():
    r = svc.interpretar("recetas_por_area", "generame las recetas de este listado", _archivos_recetas(),
                        ejecutar_llm=_stub("[]"))
    assert r["cantidad_archivos"] == 54
    assert r["advertencias"] == []
    assert r["programa"] == [{"op": "expandir_por_catalogo", "args": {}}]
    assert r["sin_indicaciones"] is True
    # el listado usa HD, que tiene el defecto real documentado en el plan
    assert any(h["regla"] == "descripcion_inconsistente_entre_variantes" for h in r["hallazgos"])


def test_recetas_ejecutar_produce_zip_con_54_archivos_agrupados_por_area(tmp_path, monkeypatch):
    import paths
    from app.ai.dsl import log as dsl_log

    monkeypatch.setattr(paths, "app_base_dir", lambda: str(tmp_path))
    programa = [{"op": "expandir_por_catalogo", "args": {}}]
    salida = svc.ejecutar("recetas_por_area", programa, _archivos_recetas(), texto_usuario="prueba de log")
    assert salida.resumen["cantidad_archivos"] == 54
    assert salida.nombre.endswith("_asistente.zip")
    with zipfile.ZipFile(io.BytesIO(salida.contenido)) as zf:
        nombres = zf.namelist()
    assert len(nombres) == 54
    assert any(n.startswith("HD/") for n in nombres)
    assert any(n.startswith("GPS2/") for n in nombres)

    ultimo = dsl_log.leer_eventos()[-1]
    assert ultimo["operacion"] == "recetas_por_area"
    assert ultimo["cantidad_archivos"] == 54


def test_recetas_programa_invalido_lanza_asistente_error():
    with pytest.raises(svc.AsistenteError):
        svc.ejecutar("recetas_por_area", [{"op": "operacion_inventada", "args": {}}], _archivos_recetas())


def test_programa_mal_formado_lanza_asistente_error():
    with pytest.raises(svc.AsistenteError):
        svc.ejecutar("recetas_por_area", "no es una lista", _archivos_recetas())


# -- plantilla ------------------------------------------------------------------


def test_plantilla_sin_indicaciones_es_igual_a_la_pantalla():
    r = svc.interpretar("plantilla", "genera los archivos", _archivos_plantilla(), ejecutar_llm=_stub("[]"))
    assert r["cantidad_archivos"] == 14
    assert r["columna_nombre_archivo"] == "sellado"
    assert [(c["linea"], c["columna"]) for c in r["campos"]] == [(1, "sellado"), (3, "amortiguador")]


def test_plantilla_indicaciones_cambian_que_columna_llena_cada_campo():
    stub = _stub(
        '[{"op": "asignar_columna", "args": {"linea": 1, "columna": "amortiguador"}},'
        ' {"op": "asignar_columna", "args": {"linea": 3, "columna": "sellado"}},'
        ' {"op": "filtrar_filas", "args": {"columna": "area", "valor": "HD"}}]'
    )
    r = svc.interpretar("plantilla", "invertilos y solo HD", _archivos_plantilla(), ejecutar_llm=stub)
    assert [(c["linea"], c["columna"]) for c in r["campos"]] == [(1, "amortiguador"), (3, "sellado")]
    assert r["columna_nombre_archivo"] == "amortiguador"
    assert r["cantidad_archivos"] < 14

    salida = svc.ejecutar("plantilla", r["programa"], _archivos_plantilla())
    with zipfile.ZipFile(io.BytesIO(salida.contenido)) as zf:
        primero = zf.namelist()[0]
        contenido = zf.read(primero).decode("utf-8")
    assert contenido.startswith(f"#Codigo;{primero.rsplit('.', 1)[0]}")


def test_plantilla_filtro_sin_filas_vuelve_como_error_de_validacion():
    stub = _stub('[{"op": "filtrar_filas", "args": {"columna": "area", "valor": "NOEXISTE"}}]')
    r = svc.interpretar("plantilla", "solo NOEXISTE", _archivos_plantilla(), ejecutar_llm=stub)
    assert "Ninguna fila" in r["error_validacion"]


def test_plantilla_misma_linea_a_dos_columnas_es_error_visible():
    stub = _stub(
        '[{"op": "asignar_columna", "args": {"linea": 1, "columna": "amortiguador"}},'
        ' {"op": "asignar_columna", "args": {"linea": 1, "columna": "sellado"}}]'
    )
    r = svc.interpretar("plantilla", "el codigo del amortiguador y del sellado", _archivos_plantilla(),
                        ejecutar_llm=stub)
    assert "dos columnas distintas" in r["error_validacion"]
