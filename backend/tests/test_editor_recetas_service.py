import io

import openpyxl
import pytest

from app.services import editor_recetas_service as svc  # noqa: F401  (side effect: agrega app/core/ a sys.path)

import paths

CSV_TEXTO = (
    "List separator=,Decimal symbol=.\r\n"
    "Recipe_1\r\n"
    "LANGID_409,P1,P2,P3,P4,P5,P6\r\n"
    "3,1,2,3,4,5,6\r\n"
    "Temperatura,45,50,55,60,65,70\r\n"
    "Presion,1.5,2.2,2.7,3.1,3.6,4.4\r\n"
    "Dureza,10,20,30,40,50,60\r\n"
)


def _csv_bytes() -> bytes:
    return CSV_TEXTO.encode("utf-8")


def _xlsx_editado(csv_bytes: bytes, nombre_parametro: str, producto: str, valor_nuevo) -> bytes:
    """Simula el paso 1 (exportar) + la edicion manual en Excel."""
    xlsx_bytes, _nombre, _resumen = svc.exportar_excel(csv_bytes, "receta.csv")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb[svc.csv_recipe.SHEET_RECETAS]
    headers = [c.value for c in ws[1]]
    col = headers.index(nombre_parametro) + 1
    fila = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == producto)
    ws.cell(row=fila, column=col, value=valor_nuevo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_exportar_excel_devuelve_xlsx_valido():
    xlsx_bytes, nombre, resumen = svc.exportar_excel(_csv_bytes(), "receta.csv")
    assert nombre == "receta (editable).xlsx"
    assert resumen == {"n_productos": 6, "n_parametros": 3}
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert svc.csv_recipe.SHEET_RECETAS in wb.sheetnames


def test_exportar_excel_archivo_invalido_lanza_error():
    with pytest.raises(svc.EditorRecetasError):
        svc.exportar_excel(b"no es una matriz de receta\n", "cualquiera.csv")


def test_previsualizar_cambios_sin_alerta():
    csv_bytes = _csv_bytes()
    xlsx_editado = _xlsx_editado(csv_bytes, "Temperatura", "P2", 99)

    reporte = svc.previsualizar_cambios(csv_bytes, "receta.csv", xlsx_editado, "receta (editable).xlsx")

    assert reporte["bloqueado"] is False
    assert reporte["cantidad_cambios"] == 1
    assert reporte["cambios"][0] == {
        "parametro": "Temperatura", "producto": "P2", "valor_anterior": "50", "valor_nuevo": "99",
    }


def test_previsualizar_cambios_con_alerta_queda_bloqueado():
    csv_bytes = _csv_bytes()
    # un salto de linea real en el valor editado corrompe la re-lectura del
    # motor (ver test_csv_recipe.py) y dispara la ALERTA de autoverificacion.
    xlsx_editado = _xlsx_editado(csv_bytes, "Temperatura", "P1", "5\n6")

    reporte = svc.previsualizar_cambios(csv_bytes, "receta.csv", xlsx_editado, "receta (editable).xlsx")

    assert reporte["bloqueado"] is True
    assert any(a.startswith("ALERTA") for a in reporte["advertencias"])


def test_aplicar_y_descargar_sin_alerta(tmp_path, monkeypatch):
    monkeypatch.setattr(paths, "app_base_dir", lambda: str(tmp_path))
    csv_bytes = _csv_bytes()
    xlsx_editado = _xlsx_editado(csv_bytes, "Dureza", "P3", 999)

    salida_bytes, nombre_salida, resumen = svc.aplicar_y_descargar(
        csv_bytes, "receta.csv", xlsx_editado, "receta (editable).xlsx"
    )

    assert resumen["cantidad_cambios"] == 1
    assert "999" in salida_bytes.decode("utf-8")
    assert nombre_salida.endswith(".csv")


def test_aplicar_y_descargar_con_alerta_rechaza(tmp_path, monkeypatch):
    monkeypatch.setattr(paths, "app_base_dir", lambda: str(tmp_path))
    csv_bytes = _csv_bytes()
    xlsx_editado = _xlsx_editado(csv_bytes, "Temperatura", "P1", "5\n6")

    with pytest.raises(svc.EditorRecetasError, match="ALERTA"):
        svc.aplicar_y_descargar(csv_bytes, "receta.csv", xlsx_editado, "receta (editable).xlsx")


def test_aplicar_y_descargar_deja_rastro_en_el_log(tmp_path, monkeypatch):
    monkeypatch.setattr(paths, "app_base_dir", lambda: str(tmp_path))
    import log_jsonl

    csv_bytes = _csv_bytes()
    xlsx_editado = _xlsx_editado(csv_bytes, "Presion", "P4", 12.5)
    svc.aplicar_y_descargar(csv_bytes, "receta.csv", xlsx_editado, "receta (editable).xlsx")

    eventos = log_jsonl.leer_eventos("editor_recetas")
    assert len(eventos) == 1
    assert eventos[0]["archivo"] == "receta.csv"
    assert eventos[0]["cantidad_cambios"] == 1
    assert eventos[0]["productos"] == 6


def test_aplicar_y_descargar_con_alerta_no_deja_rastro_en_el_log(tmp_path, monkeypatch):
    monkeypatch.setattr(paths, "app_base_dir", lambda: str(tmp_path))
    import log_jsonl

    csv_bytes = _csv_bytes()
    xlsx_editado = _xlsx_editado(csv_bytes, "Temperatura", "P1", "5\n6")
    with pytest.raises(svc.EditorRecetasError):
        svc.aplicar_y_descargar(csv_bytes, "receta.csv", xlsx_editado, "receta (editable).xlsx")

    assert log_jsonl.leer_eventos("editor_recetas") == []
